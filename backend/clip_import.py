"""
Lector de exports del panel de Clip (ventas y catalogo).

Modulo puro: no toca Mongo, no importa FastAPI, no pide credenciales. Sirve para
convertir un archivo que un humano descargo del panel de Clip en filas
normalizadas, listas para cargarse como ventas de Casa Dorelia.

Convenciones reales del panel de Clip, verificadas contra
`Catalog_20260928025118.xlsx` (export de SJI del 28/09, adjunto en BOS-66):

- El archivo es .xlsx con dos hojas: `Ayuda` (instrucciones, sin datos) y la
  hoja de datos (`Productos` en el export de catalogo).
- Los encabezados van en la fila 1 y **la fila 2 es texto de instrucciones**,
  no un renglon de datos ("Obligatorio. Maximo de 2...", etc.).
- Los encabezados estan en espanol y con acentos.

Como el export de **ventas** todavia no lo tenemos a la vista, el mapeo de
columnas es por sinonimos y, si no reconoce las columnas minimas, el modulo
**falla diciendo exactamente que encabezados vio** (`describe_table`) en lugar
de adivinar. Un solo `inspect` sobre el primer archivo real nos da el esquema.

Uso rapido:
    python clip_import.py inspect "Ventas_2026100100000.xlsx"
    python clip_import.py sales "ventas.csv" --json
    python clip_import.py catalog "Catalog_20260928025118.xlsx"
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import unicodedata
from dataclasses import dataclass, field, asdict
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

# Misma tasa que usa `create_sale` en server.py. El monto que cobra Clip ya
# incluye IVA, asi que al importar hay que desglosar, no volver a sumar.
IVA_RATE = 0.16

# Hojas del libro que nunca traen datos.
IGNORED_SHEETS = {"ayuda", "help", "instrucciones"}

# Estados de transaccion que cuentan como venta cobrada.
PAID_STATUSES = {
    "aprobada", "aprobado", "approved", "completada", "completado", "completed",
    "pagada", "pagado", "paid", "exitosa", "exitoso", "success", "successful",
    "liquidada", "cobrada", "cobrado",
}
# Estados que explicitamente NO son venta (devoluciones y rechazos).
REVERSED_STATUSES = {
    "devolucion", "devuelta", "devuelto", "reembolsada", "reembolsado",
    "refunded", "refund", "cancelada", "cancelado", "canceled", "cancelled",
    "rechazada", "rechazado", "declined", "denied", "fallida", "failed",
    "expirada", "expirado", "chargeback", "contracargo",
}

# Sinonimos de encabezado -> campo canonico. Las llaves van normalizadas
# (sin acentos, minusculas, sin signos).
SALES_HEADER_SYNONYMS: Dict[str, str] = {
    # identificador de la transaccion en Clip
    "id de transaccion": "transaction_id",
    "id transaccion": "transaction_id",
    "id de la transaccion": "transaction_id",
    "transaction id": "transaction_id",
    "id de pago": "transaction_id",
    "id de venta": "transaction_id",
    "id": "transaction_id",
    # folio / recibo / referencia
    "folio": "receipt_no",
    "recibo": "receipt_no",
    "no de recibo": "receipt_no",
    "numero de recibo": "receipt_no",
    "receipt": "receipt_no",
    "receipt no": "receipt_no",
    "referencia": "receipt_no",
    "ticket": "receipt_no",
    # fecha y hora
    "fecha": "date",
    "fecha y hora": "date",
    "fecha de la transaccion": "date",
    "fecha de transaccion": "date",
    "fecha de pago": "date",
    "date": "date",
    "datetime": "date",
    "hora": "time",
    "time": "time",
    # dinero
    "monto": "amount",
    "monto total": "amount",
    "monto de la venta": "amount",
    "importe": "amount",
    "total": "amount",
    "total cobrado": "amount",
    "amount": "amount",
    "venta": "amount",
    "propina": "tip",
    "tip": "tip",
    "comision": "fee",
    "comision clip": "fee",
    "fee": "fee",
    "monto neto": "net_amount",
    "neto": "net_amount",
    "net": "net_amount",
    # medio de pago y estado
    "metodo de pago": "payment_method",
    "forma de pago": "payment_method",
    "tipo de pago": "payment_method",
    "medio de pago": "payment_method",
    "payment method": "payment_method",
    "estado": "status",
    "estatus": "status",
    "status": "status",
    # origen
    "sucursal": "branch",
    "terminal": "branch",
    "punto de venta": "branch",
    "branch": "branch",
    "cajero": "cashier",
    "usuario": "cashier",
    # detalle de producto (solo si el export lo trae)
    "producto": "product_name",
    "nombre del producto": "product_name",
    "articulo": "product_name",
    "concepto": "product_name",
    "descripcion": "product_name",
    "cantidad": "quantity",
    "piezas": "quantity",
    "unidades": "quantity",
    "quantity": "quantity",
    "precio unitario": "unit_price",
    "precio": "unit_price",
    "unit price": "unit_price",
}

CATALOG_HEADER_SYNONYMS: Dict[str, str] = {
    "id de producto": "clip_product_id",
    "nombre del producto": "name",
    "descripcion del producto": "description",
    "categoria": "category",
    "nombre de la opcion 1": "option_1_name",
    "nombre de la variante 1": "variant_1_name",
    "nombre de la opcion 2": "option_2_name",
    "nombre de la variante 2": "variant_2_name",
    "precio": "price",
    "inventario": "inventory",
    "alerta de minimo": "min_alert",
    "codigo de barras": "barcode",
}


class ClipImportError(ValueError):
    """El archivo no se parece a un export de Clip utilizable."""


# --------------------------------------------------------------------------
# Normalizacion de texto, dinero y fechas
# --------------------------------------------------------------------------

def normalize_header(value: Any) -> str:
    """'ID de Transacción #1' -> 'id de transaccion 1'."""
    if value is None:
        return ""
    text = unicodedata.normalize("NFKD", str(value))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.replace("#", " ")
    text = re.sub(r"[^0-9a-zA-Z]+", " ", text)
    return re.sub(r"\s+", " ", text).strip().lower()


def parse_money(value: Any) -> Optional[float]:
    """Acepta 1234.5, '$1,234.50', '1.234,50', '(45.00)' y '' -> None."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text:
        return None
    negative = text.startswith("(") and text.endswith(")")
    text = text.strip("()")
    text = re.sub(r"[^0-9,.\-]", "", text)
    if not text or text in {"-", ".", ","}:
        return None
    if "," in text and "." in text:
        # El separador que aparece al final es el decimal.
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        # Una sola coma con 1 o 2 decimales es decimal; si no, es de miles.
        tail = text.split(",")[-1]
        text = text.replace(",", "." if len(tail) in (1, 2) else "")
    try:
        amount = float(text)
    except ValueError:
        return None
    return -amount if negative else amount


_DATE_FORMATS = (
    "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d",
    "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y",
    "%d-%m-%Y %H:%M:%S", "%d-%m-%Y %H:%M", "%d-%m-%Y",
    "%m/%d/%Y %H:%M:%S", "%m/%d/%Y %H:%M",
    "%d/%m/%y %H:%M", "%d/%m/%y",
)


def parse_datetime(value: Any, time_value: Any = None) -> Optional[datetime]:
    """Convierte fecha (y hora opcional en otra columna) a datetime naive."""
    if value is None:
        return None
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        parsed = datetime(value.year, value.month, value.day)
    else:
        text = str(value).strip().replace("Z", "")
        if not text:
            return None
        parsed = None
        for fmt in _DATE_FORMATS:
            try:
                parsed = datetime.strptime(text, fmt)
                break
            except ValueError:
                continue
        if parsed is None:
            try:
                parsed = datetime.fromisoformat(text)
            except ValueError:
                return None
    if time_value is not None and parsed.hour == 0 and parsed.minute == 0:
        hhmm = _parse_time(time_value)
        if hhmm:
            parsed = parsed.replace(hour=hhmm[0], minute=hhmm[1], second=hhmm[2])
    return parsed


def _parse_time(value: Any) -> Optional[Tuple[int, int, int]]:
    if isinstance(value, datetime):
        return value.hour, value.minute, value.second
    text = str(value).strip().lower()
    if not text:
        return None
    pm = "p" in text
    am = "a" in text
    match = re.search(r"(\d{1,2}):(\d{2})(?::(\d{2}))?", text)
    if not match:
        return None
    hour, minute = int(match.group(1)), int(match.group(2))
    second = int(match.group(3) or 0)
    if pm and hour < 12:
        hour += 12
    if am and hour == 12:
        hour = 0
    if not (0 <= hour <= 23 and 0 <= minute <= 59):
        return None
    return hour, minute, second


def normalize_payment_method(value: Any) -> str:
    """Devuelve el vocabulario que ya usa `SaleBase.payment_method`."""
    text = normalize_header(value)
    if not text:
        return "desconocido"
    if "efectivo" in text or "cash" in text:
        return "efectivo"
    if "transferencia" in text or "spei" in text:
        return "transferencia"
    if any(token in text for token in (
        "tarjeta", "card", "credito", "credit", "debito", "debit",
        "visa", "mastercard", "amex", "american express", "contactless", "nfc",
    )):
        return "tarjeta"
    if "link" in text or "qr" in text or "codi" in text:
        return "tarjeta"
    return text


def classify_status(value: Any) -> str:
    """-> 'paid' | 'reversed' | 'unknown'. Sin columna de estado asume cobrada."""
    text = normalize_header(value)
    if not text:
        return "paid"
    for token in text.split():
        if token in REVERSED_STATUSES:
            return "reversed"
        if token in PAID_STATUSES:
            return "paid"
    if text in PAID_STATUSES:
        return "paid"
    if text in REVERSED_STATUSES:
        return "reversed"
    return "unknown"


# --------------------------------------------------------------------------
# Lectura cruda del archivo
# --------------------------------------------------------------------------

@dataclass
class RawTable:
    """Encabezados + filas crudas de la hoja de datos."""
    source: str
    sheet: Optional[str]
    headers: List[str]
    rows: List[List[Any]]
    skipped_instruction_row: bool = False

    def as_dicts(self) -> List[Dict[str, Any]]:
        out = []
        for row in self.rows:
            padded = list(row) + [None] * (len(self.headers) - len(row))
            out.append({h: padded[i] for i, h in enumerate(self.headers)})
        return out


def _is_instruction_row(row: Sequence[Any]) -> bool:
    """La fila 2 del template de Clip es ayuda, no datos.

    Se reconoce porque es texto largo y empieza con las frases del template
    ('Obligatorio.', 'En caso de', 'Puedes escribir', 'Maximo de ...').
    """
    cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
    if not cells:
        return False
    hints = ("obligatorio", "en caso de", "puedes escribir", "maximo de",
             "escribir el numero", "codigo unico", "opcional")
    hit = sum(1 for c in cells if normalize_header(c).startswith(hints))
    return hit >= max(2, len(cells) // 2)


def read_table(path: str | Path, sheet: Optional[str] = None) -> RawTable:
    """Lee .xlsx o .csv y devuelve la tabla de datos ya sin filas de ayuda."""
    path = Path(path)
    if not path.exists():
        raise ClipImportError(f"No existe el archivo: {path}")
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        return _read_xlsx(path, sheet)
    if suffix in (".csv", ".txt", ".tsv"):
        return _read_csv(path)
    raise ClipImportError(
        f"Formato no soportado: '{suffix}'. Exporta el reporte como CSV o XLSX "
        f"desde el panel de Clip."
    )


def _read_xlsx(path: Path, sheet: Optional[str]) -> RawTable:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - openpyxl esta en requirements
        raise ClipImportError("Falta openpyxl para leer .xlsx") from exc

    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        if sheet:
            if sheet not in workbook.sheetnames:
                raise ClipImportError(
                    f"La hoja '{sheet}' no existe. Hojas: {workbook.sheetnames}"
                )
            candidates = [workbook[sheet]]
        else:
            candidates = [
                ws for ws in workbook.worksheets
                if normalize_header(ws.title) not in IGNORED_SHEETS
            ] or list(workbook.worksheets)

        best: Optional[RawTable] = None
        for worksheet in candidates:
            table = _table_from_rows(
                path.name, worksheet.title,
                [list(r) for r in worksheet.iter_rows(values_only=True)],
            )
            if best is None or len(table.rows) > len(best.rows):
                best = table
        if best is None:
            raise ClipImportError(f"El libro '{path.name}' no tiene hojas.")
        return best
    finally:
        workbook.close()


def _read_csv(path: Path) -> RawTable:
    for encoding in ("utf-8-sig", "latin-1"):
        try:
            text = path.read_text(encoding=encoding)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover - latin-1 nunca falla
        raise ClipImportError(f"No pude decodificar {path.name}")
    sample = "\n".join(text.splitlines()[:5])
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","
    rows = [list(r) for r in csv.reader(text.splitlines(), delimiter=delimiter)]
    return _table_from_rows(path.name, None, rows)


def _table_from_rows(source: str, sheet: Optional[str],
                     rows: List[List[Any]]) -> RawTable:
    # Salta renglones vacios de arriba (los exports suelen traer titulo/logo).
    index = 0
    while index < len(rows) and not any(
        c is not None and str(c).strip() for c in rows[index]
    ):
        index += 1
    if index >= len(rows):
        return RawTable(source, sheet, [], [])

    headers = [str(c).strip() if c is not None else "" for c in rows[index]]
    while headers and headers[-1] == "":
        headers.pop()
    body = rows[index + 1:]

    skipped = False
    if body and _is_instruction_row(body[0]):
        body = body[1:]
        skipped = True

    data = [
        list(r) for r in body
        if any(c is not None and str(c).strip() for c in r)
    ]
    return RawTable(source, sheet, headers, data, skipped)


def map_columns(headers: Sequence[str],
                synonyms: Dict[str, str]) -> Tuple[Dict[str, int], List[str]]:
    """-> ({campo: indice}, [encabezados no reconocidos])."""
    mapping: Dict[str, int] = {}
    unknown: List[str] = []
    for position, header in enumerate(headers):
        key = normalize_header(header)
        if not key:
            continue
        field_name = synonyms.get(key)
        if field_name is None:
            unknown.append(header)
            continue
        mapping.setdefault(field_name, position)
    return mapping, unknown


def describe_table(path: str | Path, sheet: Optional[str] = None) -> Dict[str, Any]:
    """Radiografia del archivo: hoja, encabezados, mapeo y 3 filas de muestra.

    Es el comando que se corre sobre el **primer** export real de ventas: deja
    por escrito el esquema de Clip sin tener que adivinarlo.
    """
    table = read_table(path, sheet)
    sales_map, sales_unknown = map_columns(table.headers, SALES_HEADER_SYNONYMS)
    catalog_map, _ = map_columns(table.headers, CATALOG_HEADER_SYNONYMS)
    return {
        "source": table.source,
        "sheet": table.sheet,
        "row_count": len(table.rows),
        "headers": table.headers,
        "skipped_instruction_row": table.skipped_instruction_row,
        "looks_like": (
            "catalog" if len(catalog_map) > len(sales_map) else "sales"
        ),
        "sales_columns_detected": sorted(sales_map),
        "sales_columns_unknown": sales_unknown,
        "catalog_columns_detected": sorted(catalog_map),
        "sample_rows": [
            {h: (v.isoformat() if isinstance(v, (datetime, date)) else v)
             for h, v in row.items()}
            for row in table.as_dicts()[:3]
        ],
    }


# --------------------------------------------------------------------------
# Ventas
# --------------------------------------------------------------------------

@dataclass
class ClipSale:
    """Una venta de Clip normalizada, lista para cargarse."""
    dedup_key: str
    occurred_at: Optional[datetime]
    gross_amount: float            # lo que pago el cliente, IVA incluido
    subtotal: float                # gross / 1.16
    tax: float                     # gross - subtotal
    payment_method: str
    status: str                    # paid | reversed | unknown
    transaction_id: Optional[str] = None
    receipt_no: Optional[str] = None
    tip: Optional[float] = None
    fee: Optional[float] = None
    branch: Optional[str] = None
    product_name: Optional[str] = None
    quantity: Optional[int] = None
    unit_price: Optional[float] = None
    source_row: int = 0
    warnings: List[str] = field(default_factory=list)

    def to_json(self) -> Dict[str, Any]:
        data = asdict(self)
        data["occurred_at"] = (
            self.occurred_at.isoformat() if self.occurred_at else None
        )
        return data


@dataclass
class SalesImport:
    """Resultado de leer un export de ventas."""
    source: str
    sheet: Optional[str]
    headers: List[str]
    unknown_headers: List[str]
    sales: List[ClipSale]
    rejected: List[Dict[str, Any]]

    @property
    def paid(self) -> List[ClipSale]:
        return [s for s in self.sales if s.status == "paid"]

    @property
    def gross_total(self) -> float:
        return round(sum(s.gross_amount for s in self.paid), 2)

    def summary(self) -> Dict[str, Any]:
        dates = [s.occurred_at for s in self.paid if s.occurred_at]
        by_method: Dict[str, float] = {}
        for sale in self.paid:
            by_method[sale.payment_method] = round(
                by_method.get(sale.payment_method, 0.0) + sale.gross_amount, 2
            )
        return {
            "source": self.source,
            "sheet": self.sheet,
            "rows_parsed": len(self.sales),
            "paid_count": len(self.paid),
            "reversed_count": sum(1 for s in self.sales if s.status == "reversed"),
            "unknown_status_count": sum(
                1 for s in self.sales if s.status == "unknown"
            ),
            "rejected_count": len(self.rejected),
            "duplicate_keys": self.duplicate_keys(),
            "gross_total": self.gross_total,
            "subtotal_total": round(sum(s.subtotal for s in self.paid), 2),
            "tax_total": round(sum(s.tax for s in self.paid), 2),
            "gross_by_payment_method": by_method,
            "first_sale_at": min(dates).isoformat() if dates else None,
            "last_sale_at": max(dates).isoformat() if dates else None,
            "unknown_headers": self.unknown_headers,
        }

    def duplicate_keys(self) -> List[str]:
        seen: Dict[str, int] = {}
        for sale in self.sales:
            seen[sale.dedup_key] = seen.get(sale.dedup_key, 0) + 1
        return sorted(k for k, n in seen.items() if n > 1)

    def deduplicated(self) -> List[ClipSale]:
        """Primera aparicion de cada `dedup_key`, solo ventas cobradas."""
        out: List[ClipSale] = []
        seen: set[str] = set()
        for sale in self.paid:
            if sale.dedup_key in seen:
                continue
            seen.add(sale.dedup_key)
            out.append(sale)
        return out


def split_tax(gross_amount: float, iva_rate: float = IVA_RATE) -> Tuple[float, float]:
    """Desglosa un monto que YA incluye IVA. Nunca vuelve a sumar impuesto."""
    subtotal = round(gross_amount / (1 + iva_rate), 2)
    return subtotal, round(gross_amount - subtotal, 2)


def build_dedup_key(transaction_id: Optional[str], receipt_no: Optional[str],
                    occurred_at: Optional[datetime], gross_amount: float,
                    branch: Optional[str]) -> str:
    """Clave estable por transaccion.

    Prefiere el id de Clip; si el export no lo trae, usa el folio; y si no hay
    ninguno, un hash de (sucursal, fecha-hora, monto) para que reimportar el
    mismo archivo no duplique ventas.
    """
    if transaction_id:
        return f"clip:{str(transaction_id).strip()}"
    if receipt_no:
        return f"folio:{str(receipt_no).strip()}"
    raw = "|".join([
        (branch or "").strip().lower(),
        occurred_at.isoformat() if occurred_at else "",
        f"{gross_amount:.2f}",
    ])
    return "hash:" + hashlib.sha1(raw.encode("utf-8")).hexdigest()[:16]


def parse_sales(path: str | Path, sheet: Optional[str] = None,
                iva_rate: float = IVA_RATE) -> SalesImport:
    """Lee un export de ventas de Clip y normaliza cada renglon."""
    table = read_table(path, sheet)
    if not table.headers:
        raise ClipImportError(f"'{table.source}' no tiene encabezados.")

    mapping, unknown = map_columns(table.headers, SALES_HEADER_SYNONYMS)
    if "amount" not in mapping:
        raise ClipImportError(
            "No encontre la columna de monto en "
            f"'{table.source}'{f' (hoja {table.sheet})' if table.sheet else ''}. "
            f"Encabezados vistos: {table.headers}. "
            "Corre `clip_import.py inspect <archivo>` y agrega el sinonimo real "
            "a SALES_HEADER_SYNONYMS."
        )

    def cell(row: Sequence[Any], name: str) -> Any:
        index = mapping.get(name)
        if index is None or index >= len(row):
            return None
        return row[index]

    sales: List[ClipSale] = []
    rejected: List[Dict[str, Any]] = []

    for offset, row in enumerate(table.rows):
        row_number = offset + 2  # fila 1 = encabezados
        gross = parse_money(cell(row, "amount"))
        if gross is None:
            rejected.append({"row": row_number, "reason": "monto ilegible o vacio"})
            continue

        warnings: List[str] = []
        occurred_at = parse_datetime(cell(row, "date"), cell(row, "time"))
        if occurred_at is None:
            warnings.append("sin fecha legible: hay que fijarla al cargar")

        status = classify_status(cell(row, "status"))
        if gross < 0 and status == "paid":
            status = "reversed"
            warnings.append("monto negativo: se trata como devolucion")

        transaction_id = _clean_text(cell(row, "transaction_id"))
        receipt_no = _clean_text(cell(row, "receipt_no"))
        branch = _clean_text(cell(row, "branch"))
        if not transaction_id and not receipt_no:
            warnings.append("sin id de Clip ni folio: dedup por hash de la fila")

        subtotal, tax = split_tax(abs(gross), iva_rate)
        quantity = _parse_int(cell(row, "quantity"))

        sales.append(ClipSale(
            dedup_key=build_dedup_key(
                transaction_id, receipt_no, occurred_at, abs(gross), branch
            ),
            occurred_at=occurred_at,
            gross_amount=round(abs(gross), 2),
            subtotal=subtotal,
            tax=tax,
            payment_method=normalize_payment_method(cell(row, "payment_method")),
            status=status,
            transaction_id=transaction_id,
            receipt_no=receipt_no,
            tip=parse_money(cell(row, "tip")),
            fee=parse_money(cell(row, "fee")),
            branch=branch,
            product_name=_clean_text(cell(row, "product_name")),
            quantity=quantity,
            unit_price=parse_money(cell(row, "unit_price")),
            source_row=row_number,
            warnings=warnings,
        ))

    return SalesImport(
        source=table.source, sheet=table.sheet, headers=table.headers,
        unknown_headers=unknown, sales=sales, rejected=rejected,
    )


def _clean_text(value: Any) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _parse_int(value: Any) -> Optional[int]:
    amount = parse_money(value)
    return int(round(amount)) if amount is not None else None


# --------------------------------------------------------------------------
# Catalogo
# --------------------------------------------------------------------------

@dataclass
class ClipCatalogRow:
    """Un renglon del export de catalogo (producto o variante de producto)."""
    clip_product_id: Optional[str]
    name: str
    category: Optional[str]
    price: Optional[float]
    option_name: Optional[str] = None
    variant_name: Optional[str] = None
    description: Optional[str] = None
    inventory: Optional[float] = None
    barcode: Optional[str] = None
    source_row: int = 0

    @property
    def display_name(self) -> str:
        """'Mocha 16 oz' — como el cliente lo ve en el ticket."""
        return f"{self.name} {self.variant_name}".strip() if self.variant_name else self.name


def parse_catalog(path: str | Path, sheet: Optional[str] = None) -> List[ClipCatalogRow]:
    """Lee el export de catalogo de Clip (hoja `Productos`).

    Ojo: la columna `Inventario` del export viene sembrada en 10000 y **no
    refleja la venta real** (SJI lo confirmo en BOS-80 el 01/10: 45 de 47
    renglones en 10000 y 2 en 9999 desde la apertura del 21/09). No usarla como
    fuente de unidades vendidas ni para conciliar contra el export de ventas.
    """
    table = read_table(path, sheet)
    mapping, _ = map_columns(table.headers, CATALOG_HEADER_SYNONYMS)
    if "name" not in mapping or "price" not in mapping:
        raise ClipImportError(
            f"'{table.source}' no parece export de catalogo de Clip. "
            f"Encabezados vistos: {table.headers}"
        )

    def cell(row: Sequence[Any], name: str) -> Any:
        index = mapping.get(name)
        if index is None or index >= len(row):
            return None
        return row[index]

    rows: List[ClipCatalogRow] = []
    for offset, row in enumerate(table.rows):
        name = _clean_text(cell(row, "name"))
        if not name:
            continue
        rows.append(ClipCatalogRow(
            clip_product_id=_clean_text(cell(row, "clip_product_id")),
            name=name,
            category=_clean_text(cell(row, "category")),
            price=parse_money(cell(row, "price")),
            option_name=_clean_text(cell(row, "option_1_name")),
            variant_name=_clean_text(cell(row, "variant_1_name")),
            description=_clean_text(cell(row, "description")),
            inventory=parse_money(cell(row, "inventory")),
            barcode=_clean_text(cell(row, "barcode")),
            source_row=offset + 2,
        ))
    return rows


def diff_catalog(previous: Iterable[ClipCatalogRow],
                 current: Iterable[ClipCatalogRow]) -> Dict[str, Any]:
    """Compara dos exports de catalogo por (producto, variante).

    Lo usa la publicacion de precios en casadorelia.com.mx: contesta "se movio
    algun precio entre dos exports" sin revisar 47 renglones a mano.
    """
    def index(rows: Iterable[ClipCatalogRow]) -> Dict[Tuple[str, str], ClipCatalogRow]:
        return {
            (r.name.strip().lower(), (r.variant_name or "").strip().lower()): r
            for r in rows
        }

    before, after = index(previous), index(current)
    price_changes = []
    for key in sorted(before.keys() & after.keys()):
        old, new = before[key].price, after[key].price
        if old != new:
            price_changes.append({
                "product": after[key].display_name,
                "category": after[key].category,
                "old_price": old,
                "new_price": new,
            })
    return {
        "unchanged": len(before.keys() & after.keys()) - len(price_changes),
        "price_changes": price_changes,
        "added": sorted(after[k].display_name for k in after.keys() - before.keys()),
        "removed": sorted(before[k].display_name for k in before.keys() - after.keys()),
        "identical": not price_changes
        and not (after.keys() ^ before.keys()),
    }


# --------------------------------------------------------------------------
# CLI
# --------------------------------------------------------------------------

def _json_default(value: Any) -> str:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value)


def main(argv: Optional[Sequence[str]] = None) -> int:
    parser = argparse.ArgumentParser(
        prog="clip_import",
        description="Lee exports del panel de Clip (ventas o catalogo).",
    )
    sub = parser.add_subparsers(dest="command", required=True)

    inspect_cmd = sub.add_parser("inspect", help="Radiografia del archivo")
    inspect_cmd.add_argument("path")
    inspect_cmd.add_argument("--sheet")

    sales_cmd = sub.add_parser("sales", help="Normaliza un export de ventas")
    sales_cmd.add_argument("path")
    sales_cmd.add_argument("--sheet")
    sales_cmd.add_argument("--json", action="store_true",
                           help="Imprime las ventas, no solo el resumen")

    catalog_cmd = sub.add_parser("catalog", help="Lee un export de catalogo")
    catalog_cmd.add_argument("path")
    catalog_cmd.add_argument("--sheet")
    catalog_cmd.add_argument("--diff", metavar="EXPORT_ANTERIOR",
                             help="Compara precios contra otro export")

    args = parser.parse_args(argv)

    try:
        if args.command == "inspect":
            print(json.dumps(describe_table(args.path, args.sheet),
                             indent=2, ensure_ascii=False, default=_json_default))
        elif args.command == "sales":
            result = parse_sales(args.path, args.sheet)
            payload: Dict[str, Any] = {"summary": result.summary()}
            if args.json:
                payload["sales"] = [s.to_json() for s in result.deduplicated()]
                payload["rejected"] = result.rejected
            print(json.dumps(payload, indent=2, ensure_ascii=False,
                             default=_json_default))
        else:
            rows = parse_catalog(args.path, args.sheet)
            if args.diff:
                print(json.dumps(
                    diff_catalog(parse_catalog(args.diff), rows),
                    indent=2, ensure_ascii=False, default=_json_default))
            else:
                categories: Dict[str, int] = {}
                for row in rows:
                    categories[row.category or "(sin categoria)"] = (
                        categories.get(row.category or "(sin categoria)", 0) + 1
                    )
                print(json.dumps({
                    "rows": len(rows),
                    "products": len({r.name for r in rows}),
                    "categories": categories,
                    "sample": [
                        {"product": r.display_name, "price": r.price}
                        for r in rows[:5]
                    ],
                }, indent=2, ensure_ascii=False, default=_json_default))
    except ClipImportError as exc:
        print(f"ERROR: {exc}")
        return 2
    return 0


if __name__ == "__main__":  # pragma: no cover
    raise SystemExit(main())
